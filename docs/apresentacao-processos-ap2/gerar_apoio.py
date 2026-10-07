"""Gera roteiro e ficha institucional. --url recebe o link público já publicado."""
import argparse
import json
from html import escape
from pathlib import Path
from collections import Counter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[1]
args=argparse.ArgumentParser()
args.add_argument('--url',default='')
opt=args.parse_args()
if opt.url and not opt.url.startswith('https://'): raise ValueError('Use uma URL HTTPS publicada.')
slides=json.loads((BASE/'conteudo.json').read_text(encoding='utf-8'))
assert sum(s['seconds'] for s in slides)==660
pdfmetrics.registerFont(TTFont('ArialTF','C:/Windows/Fonts/arial.ttf'))
pdfmetrics.registerFont(TTFont('ArialTFBold','C:/Windows/Fonts/arialbd.ttf'))
pdfmetrics.registerFontFamily('ArialTF',normal='ArialTF',bold='ArialTFBold')
body=ParagraphStyle('body',fontName='ArialTF',fontSize=10.5,leading=14.7,spaceAfter=8)
heading=ParagraphStyle('heading',parent=body,fontName='ArialTFBold',fontSize=13,leading=17,spaceBefore=9,spaceAfter=8,keepWithNext=True)
title=ParagraphStyle('title',parent=heading,fontSize=19,leading=24,spaceAfter=15)
small=ParagraphStyle('small',parent=body,fontSize=8.5,leading=11.5,textColor=colors.HexColor('#425665'))
cell=ParagraphStyle('cell',parent=body,fontSize=9,leading=12,spaceAfter=0)
def p(t,st=body):return Paragraph(escape(t).replace('\n','<br/>'),st)
def header(canvas,doc):
 canvas.saveState();canvas.setFont('ArialTF',9)
 canvas.drawString(48,A4[1]-28,'ULBRA — Universidade Luterana do Brasil')
 canvas.drawString(48,A4[1]-42,'Disciplina: Processos de Engenharia de Software')
 canvas.setFont('ArialTF',8);canvas.drawRightString(A4[0]-48,26,str(doc.page));canvas.restoreState()
def build(name,story):
 SimpleDocTemplate(str(BASE/name),pagesize=A4,leftMargin=48,rightMargin=48,topMargin=67,bottomMargin=45,title=name.replace('.pdf',''),author='Arthur Scharfenberger e Lucas Oliveira da Silva').build(story,onFirstPage=header,onLaterPages=header)
def table(rows,widths):
 t=Table([[p(c,cell) for c in row] for row in rows],colWidths=widths,repeatRows=1)
 t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e5eeee')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor('#cdd7dd')),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
 return t
def clock(sec):return f'{sec//60:02}:{sec%60:02}'
qa=[
 ('Por que Scrum combina com a A2?','A2 define o ciclo iterativo e incremental. A7 concretiza ciclos curtos, objetivo de Sprint, trabalho selecionado, revisão e retrospectiva. I1/I2 recortam valor e podem ocupar mais de uma Sprint. A cadência semanal foi planejada; não afirmar que Sprints já ocorreram.'),
 ('Com só dois integrantes, quem revisa?','Arthur revisa Lucas e Lucas revisa Arthur. Focos técnicos não criam áreas exclusivas. Arthur acumula Scrum Master e desenvolvimento; o PO trata das prioridades e validação de negócio. Agenda e disponibilidade do PO precisam ser combinadas.'),
 ('Pronto significa aprovado pelo PO?','A DoD registra a conclusão técnica: critérios, revisão, testes e documentação. A validação de negócio observa o fluxo demonstrado e gera aceite ou ajustes separados. Falta de feedback não equivale a aprovação.'),
 ('Por que os incrementos têm valor?','I1 permite receber e encontrar um atendimento completo; I2 permite acompanhar e concluir. Cada fatia inclui interface, regra, armazenamento e verificação. Entregar uma classe isolada não demonstra valor ao atendente.'),
 ('Como chegaram a P, M e G?','A8R04 é a referência P. Cadastros e relações são M; autorização e desempenho são G pela incerteza e verificação transversal. Não há equivalência fixa em horas. A tabela ainda é proposta até a dupla registrar suas estimativas e o argumento do consenso.'),
 ('De onde vêm os riscos 9, 6 e 4?','Probabilidade e impacto variam de 1 a 3; multiplicamos. Escopo e acesso: 3×3=9; histórico: 2×3=6; desempenho: 2×2=4. Alto significa 6–9. O login local é evidência atual para o risco de acesso.'),
 ('Por que usar fluxograma e wireframe?','O fluxograma mostra sequência, condições e erros. O wireframe mostra seleção, campos e mensagens. Juntos, permitem validar comportamento antes da integração. Cobrem cadastro e abertura; a tela específica do técnico ainda precisa ser detalhada.'),
 ('Mostrem a origem de RF2.','Abrir a A3, p. 2, §3 e §4.1: RF03, RF04 e RF05. Seguir A4 RE2 → RF2/RN1 e HU1/HU2 → A8R03/06/07 → A9 §2.2/2.3 → C1/C2/C3/C4. A3 documenta simulação com Marcos Almeida. Não confundir RF03 da A3 com RF3 da A4/A5.'),
 ('A arquitetura exige microsserviços?','Não. São responsabilidades conceituais. API/regras, autorização e persistência podem estar na mesma aplicação. A interface chama a API; a API verifica permissão e acessa persistência. As respostas retornam a quem solicitou.'),
 ('O que acontece se gravar a OS e falhar o histórico?','A proposta é gravar os dois na mesma transação: confirmar ambos ou desfazer ambos. A interface preserva dados sem exibir sucesso falso. Reenvio deve reutilizar a identificação da operação para evitar duplicação.'),
 ('Por que o login do protótipo não resolve RNF2?','Dados e login locais no navegador não comprovam controle no servidor. A proposta verifica perfil em cada operação protegida e testa chamada direta à API, inclusive leitura ou alteração indevida de dados administrativos e valores.'),
 ('Como comprovar RNF1?','Medir operações principais com base de até 10.000 OS e registrar ambiente/cenário. Pelo menos 95% devem levar até 2 segundos. Dez usuários concorrentes e cem operações por fluxo são parâmetros propostos, ainda a confirmar; não são resultados já obtidos.'),
 ('O que mudou com o feedback?','A8 retirou itens sem requisito aprovado, explicitou tamanhos/riscos e preservou 10.000 OS. A10 ganhou diagrama e matriz com oito requisitos, vínculos com A7 e testes derivados. A3 fornecida permitiu completar a origem por página, seção e ID.'),
 ('O que está implementado e o que fica para AS?','Há interface com armazenamento local e classes Java executadas separadamente. API, autorização no servidor, banco e integração descritos na arquitetura são propostos. AS deve integrar, testar regras/falhas/desempenho e registrar evidências.'),
 ('Marcos e Chico são a mesma pessoa?','Não. Marcos Almeida é o atendente da entrevista simulada documentada na A3. Chico Mosca foi indicado como PO no planejamento A7 a partir do feedback. Não atribuir a Chico respostas da simulação nem afirmar validação real não registrada.'),
]
story=[p('AP2 • Roteiro de apresentação',title),p('Arthur Scharfenberger e Lucas Oliveira da Silva • 07/10/2026'),p('Meta: 11 minutos de exposição e 4 minutos para perguntas. Os slides 14–16 são apoio e só entram se necessários. As falas completas também estão nas notas do PowerPoint.')]
rows=[['Slide / assunto','Quem','Intervalo']];elapsed=0
md=['# Roteiro AP2 — Processos de Engenharia de Software','', 'Meta: 11 minutos de exposição + 4 minutos para perguntas. Slides 14–16 apenas como apoio.','']
for i,s in enumerate(slides[:13],1):
 interval=f'{clock(elapsed)}–{clock(elapsed+s["seconds"])}';elapsed+=s['seconds']
 rows.append([f'{i:02} • {s["title"]}',s['speaker'],interval])
story.append(table(rows,[315,95,89]))
story.extend([Spacer(1,13),p('Antes de entrar na sala',heading),p('Abram o PPTX e o PDF de reserva. Ensaiem uma vez com cronômetro, apontando os elementos dos diagramas. Alternem as respostas no treino: qualquer integrante pode ser questionado. Confiram as propostas de estimativas, o papel do PO e os arquivos A5/A9. O link público precisa ser testado em janela anônima antes de entrar na ficha final.')])
for group in [(0,3),(3,6),(6,9),(9,13)]:
 story.append(PageBreak())
 for i in range(*group):
  s=slides[i]
  h=f'{i+1:02} • {s["title"]} — {s["speaker"]} ({s["seconds"]} s)'
  story.extend([p(h,heading),p(s['notes']),p('Fonte: '+s['source'],small)])
  md += ['## '+h,'',s['notes'],'','Fonte: '+s['source'],'']
story.append(PageBreak());story.append(p('Perguntas para o ensaio',title))
md+=['## Perguntas para o ensaio','']
for q,a in qa:
 story.extend([p(q,heading),p(a)]);md.extend(['### '+q,'',a,''])
build('Roteiro-Arthur-Lucas.pdf',story)
(BASE/'Roteiro-Arthur-Lucas.md').write_text('\n'.join(md),encoding='utf-8')

urltext=opt.url or 'PENDENTE — inserir o link público da apresentação após o compartilhamento.'
sections=[
 ('1. Identificação',[
 'Atividade (código e nome): AP2 — Apresentação da segunda etapa do projeto',
 'Etapa (AP1 / AP2 / AS): AP2 | Equipe: TechFix',
 'Integrantes: Arthur Scharfenberger e Lucas Oliveira da Silva',
 'Data: 07/10/2026 | Momento: Encontro 11 | Modalidade: apresentação oral em equipe']),
 ('2. Conteúdo da atividade',[
 'Link da apresentação: '+urltext,
 'Arquivo: TechFix-AP2-Processos.pptx; cópia de segurança em PDF. São 13 slides principais e 3 de apoio, com falas distribuídas entre os dois integrantes. Tempo planejado: 11 minutos de exposição e 4 minutos de perguntas.',
 'Conteúdo: A7 — processo, papéis e fluxo; A8 — backlog, estimativas, incrementos e riscos; A9 — fluxograma e wireframe de cadastro/abertura; A10 — arquitetura, matriz e caminho A3 → requisito → backlog → modelo → componente; fechamento com qualidade, testes e evolução na AS.',
 'Rastreabilidade demonstrada: A3, p. 2, §3/§4.1 RF03–RF05 → A4 RE2/RF2/RN1 e HU1/HU2 → A5 critérios → A8R03/06/07 → A9 §2.2/§2.3 → C1–C4. A origem é uma entrevista simulada, conforme o próprio documento A3.']),
 ('3. Decisões e pendências',[
 'Decisões: usar a abertura de OS como exemplo contínuo; manter o ciclo iterativo e incremental; distinguir arquitetura proposta de implementação; reservar quatro minutos para perguntas.',
 ('Link informado para a ficha; conferir acesso público antes do envio. ' if opt.url else 'Esta ficha é um rascunho até receber um link público válido. ')+ 'A dupla ainda deve registrar consenso de estimativas, confirmar agenda do PO, conciliar a A5 avaliada com a cópia local e conferir os modelos da A9 enviada. Revisão e assinaturas permanecem sob responsabilidade dos integrantes.']),
 ('4. Declaração de uso de Inteligência Artificial',[
 '(X) A equipe utilizou ferramentas de IA nesta atividade.',
 'Data e ferramenta: 07/10/2026 — ChatGPT/Codex.',
 'Objetivo e resumo da interação: preparar a apresentação AP2 a partir do enunciado e dos artefatos A2–A10, com slides, roteiro dividido, rastreabilidade verificável e ficha padrão.',
 'Resultado aproveitado/modificado: organização do conteúdo, diagramas e tabelas editáveis, notas de fala, roteiro e ficha. A origem A3 foi conferida no arquivo fornecido; não foram criadas entrevistas, reuniões, consenso, assinaturas ou resultados de testes.',
 'Forma de verificação: confronto dos IDs e critérios com os documentos locais; inspeção dos slides e PDFs; revisão final do conteúdo e ensaio pela dupla ainda a realizar. Uso registrado no DEVLOG.md.',
 'Responsáveis pela decisão final sobre o uso: Arthur Scharfenberger e Lucas Oliveira da Silva.']),
 ('5. Responsabilidade e autoria',[
 'A equipe declara que este documento foi lido, compreendido e validado por todos os integrantes, que respondem integralmente pela correção, coerência, legalidade e autoria do conteúdo entregue. Qualquer integrante poderá ser questionado, em apresentação, sobre partes produzidas com apoio de IA.',
 'Assinatura (nomes por extenso), após revisão:',
 'Arthur Scharfenberger: ___________________________________________',
 'Lucas Oliveira da Silva: ___________________________________________'])]
stem='Ficha-Entrega-AP2' if opt.url else 'Ficha-Entrega-AP2-Rascunho'
story=[p('Entrega de atividade — AP2',title)]
for i,(h,paras) in enumerate(sections):
 if i==3:story.append(PageBreak())
 story.append(p(h,heading))
 for t in paras:
  if opt.url and t.startswith('Link da apresentação:'):
   story.append(Paragraph('Link da apresentação: <link href="'+escape(opt.url,quote=True)+'" color="#006f6f">'+escape(opt.url)+'</link>',body))
  else:story.append(p(t))
build(stem+'.pdf',story)
doc=Document(ROOT/'docs/entregas-processos/Documento_Padrao_Entrega_Atividades.docx')
for child in list(doc.element.body):
 if child.tag!=qn('w:sectPr'):doc.element.body.remove(child)
doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10.5)
doc.add_heading('Entrega de atividade — AP2',0)
for h,paras in sections:
 doc.add_heading(h,1)
 for t in paras:
  if opt.url and t.startswith('Link da apresentação:'):
   par=doc.add_paragraph('Link da apresentação: ')
   link=OxmlElement('w:hyperlink');link.set(qn('r:id'),par.part.relate_to(opt.url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
   run=OxmlElement('w:r');tx=OxmlElement('w:t');tx.text=opt.url;run.append(tx);link.append(run);par._p.append(link)
  else:doc.add_paragraph(t)
doc.save(BASE/(stem+'.docx'))
(BASE/(stem+'.md')).write_text('# Entrega de atividade — AP2\n\nDisciplina: Processos de Engenharia de Software\n\n'+'\n\n'.join('## '+h+'\n\n'+'\n\n'.join(ps) for h,ps in sections),encoding='utf-8')
print('Roteiro gerado; ficha: '+stem+'; exposição: '+str(elapsed)+' segundos.')
